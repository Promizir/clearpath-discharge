"""Generate submission visuals. Pillow and python-pptx are optional build tools."""
from pathlib import Path
from html import escape
import base64

ROOT = Path(__file__).resolve().parent
NAVY = "152C3B"
TEAL = "087F72"
IVORY = "F5F6F3"
AMBER = "D59B41"
AMD_SCREENSHOT = ROOT / "screenshots" / "04-amd-review-findings.png"
DEMO_SCREENSHOT = AMD_SCREENSHOT if AMD_SCREENSHOT.exists() else ROOT / "screenshots" / "02-review-findings.png"
DEMO_TITLE = "A real workflow.\nVisible findings."
DEMO_CAPTION = "Actual AMD-assisted run · synthetic patient" if AMD_SCREENSHOT.exists() else "Actual running dashboard · synthetic patient"
DEMO_POINTS = [
    "AMD-assisted draft reviewed before independent checks.",
    "Deterministic MCP tools return pharmacy and billing findings.",
    "Human-review packet. No autonomous discharge clearance.",
]

SLIDES = [
    ("01 / TRACK 2 — HEALTH AND WELLBEING", "A clearer path\nto discharge.", "ClearPath — AMD-assisted discharge coordination", [
        "One reviewable workflow for drafts, approvals, and findings.",
        "Synthetic-data prototype · Human approval remains explicit."]),
    ("02 / THE PROBLEM", "A change should not\nleave approval behind.", "Built for the discharge coordinator", [
        "Clinical documentation, pharmacy, and billing depend on shared information.",
        "A corrected draft can make earlier approvals and findings outdated.",
        "Coordinators need to see what is current and what still needs review."]),
    ("03 / THE WORKFLOW", "Review once.\nCheck together.", "An explicit sequence with visible outcomes", [
        "1  Generate a draft with DeepSeek-V4-Flash through the configured AMD API.",
        "2  Approve the specific revision under review.",
        "3  Run independent pharmacy and billing checks concurrently.",
        "4  Return findings for human review; persist the workflow."]),
    ("04 / THE DEMO", "Two checks.\nOne shared record.", "Structured findings, not automatic clearance", [
        "Pharmacy: the exact demo rule flags Amoxicillin / Penicillin.",
        "Billing: usage comparison identifies 2 gauze units and 1 stapler not billed.",
        "A completed check can contain findings. A failed check stays distinct.",
        "No bill is changed and no clinical decision is automated."]),
    ("05 / THE DESIGN", "Every result belongs\nto a revision.", "Traceable state changes are the core of the prototype", [
        "Edits invalidate the previous approval and stored findings.",
        "Revision checks reject results that became stale while workers ran.",
        "SQLite restores saved snapshots after restart.",
        "Interrupted work becomes a visible failure and can be retried."]),
    ("06 / THE BUILD", "AMD drafts.\nRules check records.", "DeepSeek-V4-Flash → human review → deterministic MCP checks", [
        "A live configured AMD endpoint call was verified using synthetic data.",
        "FastAPI + Pydantic: validated requests and structured results.",
        "FastMCP + asyncio: concurrent pharmacy and billing rule checks.",
        "SQLite: persist the latest workflow snapshot and its revision."]),
    ("07 / THE NEXT STEP", "Coordinate first.\nValidate before scale.", "A working foundation with an honest scope", [
        "Today: AMD-assisted drafting, synthetic records, simulated approval.",
        "Next: authenticated reviewers and validated EHR / insurer integrations.",
        "Next: clinical evaluation, multimodal inputs, and a complete audit history.",
        "A successful endpoint call demonstrates integration, not clinical validation."]),
]

def build_svg():
    svg = '''<svg xmlns="http://www.w3.org/2000/svg" width="1600" height="900" viewBox="0 0 1600 900">
<rect width="1600" height="900" fill="#f5f6f3"/>
<rect x="96" y="92" width="44" height="8" rx="4" fill="#087f72"/>
<text x="160" y="109" fill="#087f72" font-family="Arial,sans-serif" font-size="26" font-weight="bold" letter-spacing="3">CLEARPATH</text>
<text x="96" y="310" fill="#152c3b" font-family="Georgia,serif" font-size="104">A clearer path</text>
<text x="96" y="430" fill="#152c3b" font-family="Georgia,serif" font-size="104">to discharge.</text>
<text x="100" y="510" fill="#52646b" font-family="Arial,sans-serif" font-size="29">Draft. Review. Coordinate.</text>
<text x="100" y="565" fill="#52646b" font-family="Arial,sans-serif" font-size="25">One workflow. Every approval current.</text>
<rect x="1050" y="170" width="430" height="530" rx="28" fill="#152c3b"/>
<text x="1090" y="227" fill="#a9d9ce" font-family="Arial,sans-serif" font-size="18" letter-spacing="2">DISCHARGE WORKFLOW</text>
<path d="M1113 294 V600" stroke="#4c7778" stroke-width="3" fill="none"/>
<g font-family="Arial,sans-serif" font-size="27" fill="#ffffff">
<circle cx="1113" cy="300" r="13" fill="#a9d9ce"/><text x="1147" y="310">Draft prepared</text>
<circle cx="1113" cy="399" r="13" fill="#a9d9ce"/><text x="1147" y="409">Revision approved</text>
<circle cx="1113" cy="498" r="13" fill="#a9d9ce"/><text x="1147" y="508">Checks complete</text>
<circle cx="1113" cy="597" r="13" fill="#d59b41"/><text x="1147" y="607">Human review</text>
</g>
<path d="M96 776 H1504" stroke="#d7ded9"/>
<text x="100" y="826" fill="#52646b" font-family="Arial,sans-serif" font-size="21">Discharge Coordination Prototype</text>
<text x="1500" y="826" text-anchor="end" fill="#52646b" font-family="Arial,sans-serif" font-size="21">Synthetic data · Human approval</text>
</svg>'''
    (ROOT / "cover.svg").write_text(svg, encoding="utf-8")

def build_html():
    sections = []
    for i, (kicker, title, subtitle, bullets) in enumerate(SLIDES):
        if i == 3 and DEMO_SCREENSHOT.exists():
            encoded = base64.b64encode(DEMO_SCREENSHOT.read_bytes()).decode("ascii")
            items = "".join(f"<li>{escape(b)}</li>" for b in DEMO_POINTS)
            sections.append(f'<section class="slide screenshot-slide" id="slide-{i+1}"><div class="kicker">{kicker}</div><div class="demo-copy"><h1>{DEMO_TITLE.replace(chr(10), "<br>")}</h1><h2>{DEMO_CAPTION}</h2><ul>{items}</ul></div><img class="dashboard" src="data:image/png;base64,{encoded}" alt="Actual ClearPath dashboard showing doctor approval and pharmacy and billing findings for a synthetic patient"><footer><span>CLEARPATH / DISCHARGE COORDINATION</span><span>{i+1:02} / {len(SLIDES):02}</span></footer></section>')
            continue
        items = "".join(f"<li>{escape(b)}</li>" for b in bullets)
        sections.append(f'<section class="slide" id="slide-{i+1}"><div class="kicker">{kicker}</div><h1>{escape(title).replace(chr(10), "<br>")}</h1><h2>{escape(subtitle)}</h2><ul>{items}</ul><footer><span>CLEARPATH / DISCHARGE COORDINATION</span><span>{i+1:02} / {len(SLIDES):02}</span></footer></section>')
    html = '''<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>ClearPath — Pitch</title><style>
*{box-sizing:border-box}body{margin:0;background:#dce3df;color:#152c3b;font-family:Arial,sans-serif}.slide{width:1280px;height:720px;background:#f5f6f3;margin:24px auto;padding:46px 66px;position:relative;overflow:hidden;box-shadow:0 10px 40px #152c3b16}.slide:before{content:"";position:absolute;right:0;top:0;width:16px;height:100%;background:#087f72}.kicker{color:#087f72;font-size:16px;font-weight:700;letter-spacing:3px}h1{font-family:Georgia,serif;font-size:64px;line-height:1.04;font-weight:400;margin:26px 0 20px}h2{font-size:23px;font-weight:400;color:#52646b;margin:0 0 24px}ul{list-style:none;margin:0;padding:0;width:1050px}li{font-size:23px;line-height:1.35;padding:9px 0 9px 28px;position:relative}li:before{content:"";height:8px;width:8px;border-radius:50%;background:#d59b41;position:absolute;left:0;top:21px}footer{position:absolute;left:66px;right:66px;bottom:31px;display:flex;justify-content:space-between;border-top:1px solid #d3ddd6;padding-top:18px;font-size:13px;letter-spacing:1px;color:#52646b}.controls{position:fixed;right:24px;bottom:18px;z-index:10;display:flex;gap:8px}button{background:#152c3b;color:white;border:0;border-radius:8px;padding:12px 18px;cursor:pointer;font-size:14px}@media(max-width:1300px){.slide{margin:12px auto;zoom:.75}}@media(max-width:990px){.slide{zoom:.55}}@page{size:13.333in 7.5in;margin:0}@media print{body{background:white}.slide{margin:0;box-shadow:none;break-after:page;zoom:1;width:13.333in;height:7.5in}.controls{display:none}}
</style></head><body>''' + "".join(sections) + '''<nav class="controls" aria-label="Slide controls"><button onclick="step(-1)">Previous</button><button onclick="step(1)">Next</button><button onclick="window.print()">Print / PDF</button></nav><script>const slides=[...document.querySelectorAll('.slide')];let current=0;function step(delta){current=Math.max(0,Math.min(slides.length-1,current+delta));slides[current].scrollIntoView({behavior:'smooth',block:'start'});}document.addEventListener('keydown',e=>{if(['ArrowRight','PageDown','ArrowLeft','PageUp'].includes(e.key)){e.preventDefault();step(['ArrowRight','PageDown'].includes(e.key)?1:-1);}});</script></body></html>'''
    html = html.replace("</style>", ".demo-copy{width:350px}.demo-copy h1{font-size:42px;margin-top:45px}.demo-copy h2{font-size:21px;line-height:1.45;margin-bottom:30px}.demo-copy ul{width:350px}.demo-copy li{font-size:21px;margin-bottom:14px}.dashboard{position:absolute;left:470px;top:56px;width:730px;height:auto;border:1px solid #d3ddd6;box-shadow:0 8px 24px #152c3b12}</style>")
    (ROOT / "pitch.html").write_text(html, encoding="utf-8")

def build_png():
    try:
        from PIL import Image, ImageDraw, ImageFont
    except ImportError:
        print("Pillow unavailable; cover.svg is ready.")
        return
    im = Image.new("RGB", (1600, 900), "#"+IVORY)
    d = ImageDraw.Draw(im)
    fonts = Path("C:/Windows/Fonts")
    def font(size, serif=False, bold=False):
        return ImageFont.truetype(str(fonts / ("georgia.ttf" if serif else "arialbd.ttf" if bold else "arial.ttf")), size)
    def txt(x,y,s,size=25,color="#52646b",serif=False,bold=False):
        d.text((x,y),s,font=font(size,serif,bold),fill=color)
    d.rounded_rectangle((96,92,140,100),radius=4,fill="#"+TEAL)
    txt(160,83,"CLEARPATH",26,"#"+TEAL,bold=True)
    txt(96,214,"A clearer path",104,"#"+NAVY,True)
    txt(96,334,"to discharge.",104,"#"+NAVY,True)
    txt(100,480,"Draft. Review. Coordinate.",29)
    txt(100,540,"One workflow. Every approval current.",25)
    d.rounded_rectangle((1050,170,1480,700),radius=28,fill="#"+NAVY)
    txt(1090,204,"DISCHARGE WORKFLOW",18,"#a9d9ce")
    d.line((1113,294,1113,600),fill="#4c7778",width=3)
    for y,label,c in [(300,"Draft prepared","#a9d9ce"),(399,"Revision approved","#a9d9ce"),(498,"Checks complete","#a9d9ce"),(597,"Human review","#"+AMBER)]:
        d.ellipse((1100,y-13,1126,y+13),fill=c)
        txt(1147,y-15,label,27,"white")
    d.line((96,776,1504,776),fill="#d7ded9",width=1)
    txt(100,806,"Discharge Coordination Prototype",21)
    label="Synthetic data · Human approval"
    d.text((1500,806),label,font=font(21),fill="#52646b",anchor="ra")
    im.save(ROOT / "cover.png")
    print("Created cover.png")

def build_pptx():
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_SHAPE
    except ImportError:
        print("python-pptx unavailable; printable pitch.html is ready.")
        return
    prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
    def box(slide,x,y,w,h,text,size,color,font="Aptos",bold=False):
        tf=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)).text_frame
        tf.word_wrap=True;tf.margin_left=0;tf.margin_top=0;tf.margin_bottom=0;tf.margin_right=0
        for n,line in enumerate(text.split("\n")):
            p=tf.paragraphs[0] if n==0 else tf.add_paragraph()
            p.text=line;p.font.name=font;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color)
    for i,(kicker,title,subtitle,bullets) in enumerate(SLIDES):
        s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(IVORY)
        bar=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(13.16),0,Inches(.18),prs.slide_height);bar.fill.solid();bar.fill.fore_color.rgb=RGBColor.from_string(TEAL);bar.line.fill.background()
        box(s,.69,.47,11,.4,kicker,13,TEAL,bold=True)
        if i == 3 and DEMO_SCREENSHOT.exists():
            box(s,.69,1.18,3.7,1.45,DEMO_TITLE,30,NAVY,"Georgia")
            box(s,.69,2.65,3.65,.8,DEMO_CAPTION,16,"52646B")
            for j,b in enumerate(DEMO_POINTS):
                dot=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(.71),Inches(3.64+j*.93),Inches(.08),Inches(.08));dot.fill.solid();dot.fill.fore_color.rgb=RGBColor.from_string(AMBER);dot.line.fill.background()
                box(s,.94,3.52+j*.93,3.5,.85,b,16,NAVY)
            s.shapes.add_picture(str(DEMO_SCREENSHOT),Inches(4.9),Inches(.6),width=Inches(7.6))
        else:
            box(s,.69,1.0,11.8,1.9,title,47,NAVY,"Georgia")
            box(s,.69,2.95,11.8,.5,subtitle,19,"52646B")
            for j,b in enumerate(bullets):
                dot=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(.71),Inches(3.69+j*.57),Inches(.08),Inches(.08));dot.fill.solid();dot.fill.fore_color.rgb=RGBColor.from_string(AMBER);dot.line.fill.background()
                box(s,.98,3.56+j*.57,11.4,.55,b,18,NAVY)
        line=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(.69),Inches(6.84),Inches(11.94),Inches(.01));line.fill.solid();line.fill.fore_color.rgb=RGBColor.from_string("D3DDD6");line.line.fill.background()
        box(s,.69,7.02,10,.25,"CLEARPATH / DISCHARGE COORDINATION",10,"52646B")
        box(s,11.9,7.02,.8,.25,f"{i+1:02} / {len(SLIDES):02}",10,"52646B")
        s.notes_slide.notes_text_frame.text=(DEMO_CAPTION+"\n"+"\n".join(DEMO_POINTS)) if i==3 and DEMO_SCREENSHOT.exists() else subtitle+"\n"+"\n".join(bullets)
        if i == 5:
            s.notes_slide.notes_text_frame.text += "\nEvidence: submission/amd-execution-evidence.json records one successful configured API response using DeepSeek-V4-Flash. Endpoint: https://developer.amd.com.cn/radeon/api/v1. This is not hardware attestation or a clinical performance evaluation."
    prs.save(ROOT / "ClearPath-pitch.pptx")
    print("Created ClearPath-pitch.pptx")

def build_pdf():
    """Create a portable slide PDF without a browser or Office dependency."""
    try:
        from PIL import Image, ImageDraw, ImageFont, JpegImagePlugin
    except ImportError:
        return
    fonts = Path("C:/Windows/Fonts")
    pages=[]
    for i,(kicker,title,subtitle,bullets) in enumerate(SLIDES):
        im=Image.new("RGB",(1920,1080),"#"+IVORY)
        d=ImageDraw.Draw(im)
        def font(size, serif=False, bold=False):
            return ImageFont.truetype(str(fonts / ("georgia.ttf" if serif else "arialbd.ttf" if bold else "arial.ttf")),size)
        d.rectangle((1896,0,1920,1080),fill="#"+TEAL)
        d.text((100,70),kicker,font=font(24,bold=True),fill="#"+TEAL)
        is_demo = i == 3 and DEMO_SCREENSHOT.exists()
        if is_demo:
            d.multiline_text((100,185),DEMO_TITLE,font=font(61,serif=True),fill="#"+NAVY,spacing=8)
            d.multiline_text((100,390),("Actual AMD-assisted run" if AMD_SCREENSHOT.exists() else "Actual running dashboard")+"\nSynthetic patient",font=font(29),fill="#52646b",spacing=8)
            screenshot = Image.open(DEMO_SCREENSHOT).convert("RGB")
            screenshot.thumbnail((1094,866),Image.Resampling.LANCZOS)
            im.paste(screenshot,(706,86))
            d.rectangle((705,85,706+screenshot.width,86+screenshot.height),outline="#d3ddd6",width=1)
        else:
            d.multiline_text((100,145),title,font=font(92,serif=True),fill="#"+NAVY,spacing=6)
            d.text((100,430),subtitle,font=font(34),fill="#52646b")
        y=525
        bullet_font = font(28 if is_demo else 32)
        for bullet in DEMO_POINTS if is_demo else bullets:
            d.ellipse((104,y+15,116,y+27),fill="#"+AMBER)
            lines=[];line=""
            for word in bullet.split():
                proposed=(line+" "+word).strip()
                if d.textlength(proposed,font=bullet_font)>(475 if is_demo else 1630) and line:
                    lines.append(line);line=word
                else:
                    line=proposed
            lines.append(line)
            d.multiline_text((145,y),"\n".join(lines),font=bullet_font,fill="#"+NAVY,spacing=7)
            y+=max(76,len(lines)*40+20)
        d.line((100,985,1820,985),fill="#d3ddd6",width=2)
        d.text((100,1010),"CLEARPATH / DISCHARGE COORDINATION",font=font(18),fill="#52646b")
        d.text((1700,1010),f"{i+1:02} / {len(SLIDES):02}",font=font(18),fill="#52646b")
        pages.append(im)
    pages[0].save(ROOT / "ClearPath-pitch.pdf",save_all=True,append_images=pages[1:],resolution=144,title="ClearPath — Discharge Coordination",author="ClearPath project")
    pages[3].save(ROOT / "slide-preview.png")
    print("Created ClearPath-pitch.pdf")

if __name__ == "__main__":
    build_svg();build_html();build_png();build_pptx();build_pdf()
